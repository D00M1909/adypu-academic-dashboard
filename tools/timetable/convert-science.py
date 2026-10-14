# Converts the School of Science timetable workbooks (one per semester) into
# the format build.php reads.
#
#   python convert-science.py "Sem I.xlsx=science-sem1-2026-09-28.json" "Sem III.xlsx=science-sem3-2026-09-28.json"
#
# One output per workbook, so each cites its own sheet; read together because
# a teacher used in one sheet can be listed only under the other's grid.
#
# Needs openpyxl. Each workbook is one sheet: days down, the times across a
# "Day/Time" row ("9:30-10:30", with "11:30-12:00" and "2:00-2:30" as breaks),
# and a cell merged across several times is one long period. A cell is one
# or more parts split by "/", each "Subject (Div-A) (INITIALS) (room)": the
# division and the initials are optional, the room is the last bracket ("316",
# "Lecture Hall-3"). Parts side by side are labs for different divisions, or
# two courses running at once (Mathematics or Chemistry). A part with no
# initials takes the teacher of the course it names from the list under the
# grid ("Mr. Name Surname (DL)"). A part with no room (Sports, Library, Doubt
# Solving, Project/Seminar) is printed and left out.
#
# The sheet does not name its building; the rooms are in the Dental building.
import json, re, sys
import openpyxl

tidy = lambda s: re.sub(r'\s+', ' ', str(s or '')).strip()
norm = lambda s: re.sub(r'[^a-z0-9]+', ' ', s.lower()).strip()

teacher = {}   # initials => full name, from every sheet's list: Sem I uses SN, listed only under Sem III
outs = []
for arg in sys.argv[1:]:
    path, dest = arg.rsplit('=', 1)
    out = {'source': re.split(r'[\\/]', path)[-1], 'building': 'Dental Building',
           'note': 'Converted by convert-science.py (W.E.F 28/09/2026). Building not stated in the sheet: the Dental '
                   'building, per the university. Periods with no room are left out.',
           'faculty': {'': ''}, 'subjects': {}, 'classes': []}
    outs.append((out, dest))
    ws = openpyxl.load_workbook(path).worksheets[0]
    merged = {(m.min_row, m.min_col): (m.max_row, m.max_col) for m in ws.merged_cells.ranges}
    cell = lambda r, c: tidy(ws.cell(r, c).value)
    sem = next(re.search(r'Semester-?\s*([IVX]+)', cell(r, 1)).group(1) for r in range(1, 15) if re.search(r'Semester-?\s*[IVX]+', cell(r, 1)))
    label = f'B.Sc Sem {sem}'
    head = next(r for r in range(1, 20) if cell(r, 1) == 'Day/Time')
    times = {}   # column => start ("9.30"), teaching columns only
    for c in range(2, ws.max_column + 1):
        m = re.match(r'(\d{1,2}):(\d{2})-(\d{1,2}):(\d{2})', cell(head, c))
        if m and not any(cell(r, c).lower().startswith('short break') for r in range(head + 1, head + 6)):
            times[c] = f'{int(m.group(1))}.{m.group(2)}'
    lst = next(r for r in range(head, ws.max_row + 1) if cell(r, 1).startswith('Sl'))
    courses = []   # (course name, initials)
    for r in range(lst + 1, ws.max_row + 1):
        name, fac = cell(r, 3), cell(r, 4)
        if not name: break
        found = re.findall(r'([^/()]+?)\s*\(([A-Z]{2,3})\)', fac)
        for full, ini in found: teacher[ini] = tidy(re.sub(r'^(Dr|Mr|Ms|Mrs|Prof)\.?\s*', lambda m: m.group(1) + '. ', full.strip()))
        courses.append((name, found[0][1] if found else ''))
    entries = []
    for r in range(head + 1, lst):
        day = cell(r, 1)
        if day not in ('Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday'): continue
        for c, start in times.items():
            text = cell(r, c)
            if not text: continue
            last = merged.get((r, c), (r, c))[1]
            hours = sum(1 for x in range(c, last + 1) if x in times)
            for part in [tidy(p) for p in re.split(r'/(?![A-Z]{2}\))(?=\s*[A-Z][a-z])', text)]:
                brackets = re.findall(r'\(([^()]*)\)', part)
                subject = tidy(re.sub(r'\(.*', '', part))
                room = brackets[-1] if brackets and re.match(r'^(\d{3}|Lecture Hall-?\s*\d+)$', brackets[-1]) else ''
                if not room:
                    print(f'  {label} {day} {start}: no room, left out: {part}', file=sys.stderr)
                    continue
                div = next((re.sub(r'Div-?\s*', '', b) for b in brackets if re.match(r'Div-?\s*[A-Z]$', b)), '')
                who = next((b for b in brackets if re.match(r'^[A-Z]{2,3}(/[A-Z]{2,3})*$', b)), '')
                if not who:
                    key = norm(re.sub(r'\bmaths\b', 'mathematics', subject, flags=re.I))
                    who = next((i for n, i in courses if norm(n).startswith(key) or key.startswith(norm(n))), '')
                room = re.sub(r'Lecture Hall-?\s*', 'Lecture Hall ', room)
                code = f'{label}|{subject}'
                out['subjects'][code] = subject
                entries.append([day, start, hours, code, room, who, bool(re.search(r'\blab\b', subject, re.I)), f'Div {div}' if div else ''])
    out['classes'].append({'class': label, 'page': 'Sheet1', 'entries': entries})
    print(f'{label}: {len(entries)} bookings', file=sys.stderr)
# "SD/JG" is two teachers: the build reads a "/" in a name as two people.
for out, dest in outs:
    for c in out['classes']:
        for e in c['entries']:
            out['faculty'][e[5]] = ' / '.join(teacher.get(w, w) for w in e[5].split('/'))
    json.dump(out, open(dest, 'w', encoding='utf-8'), indent=1, ensure_ascii=False)
