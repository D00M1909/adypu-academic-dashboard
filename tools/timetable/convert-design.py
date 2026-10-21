# Converts one School of Design timetable workbook into the format build.php
# reads.
#
#   python convert-design.py "School of Design - UI-UX Design Timetable 2026-27.xlsx" design-uiux-2026-10-14.json
#
# Needs openpyxl. Design does not timetable by the week. Each tab is one class's
# semester calendar: a row per week (Mor. and Aft. rows), Monday to Friday
# across, each day a date column and a subject column, and one subject filling
# the whole day (9:30 to 12:30 and 1:00 to 4:00) in the class's own room, named
# once in the header ("Class Room No. 108 (First Floor)"). A subject cell merged
# over both rows is the whole day; Interior Design fills the two separately.
# A tab whose header names no room is printed and left out, except Interior
# Design's, which goes in its own room, ROOMLESS (the user's call, 21 Oct 2026).
#
# So each tab becomes a calendar, not bookings: its rooms, its subjects, and
# for every date the subject of each session, which the room page turns into
# that week's bookings. Holidays, vacation, exams, events and days away from
# the room (study tour, overseas exchange) are left out, so those days read
# as free. Each session's cell goes with it, for the source panel.
#
# A cell names its subject in full ("Physical Model Making"), by the course's
# initials ("DAF" is Designing Automotive Forms, "ESI: Module 01: ..." is
# Elementary Sketching and Illustrations), or as neither. The course list under
# the grid gives each course's name and its lead, who is the booking's teacher.
# A course in TEAM_TAUGHT is taught by its whole team, lead included.
import datetime, json, re, sys
import openpyxl

MONTHS = {m: i for i, m in enumerate(['january', 'february', 'march', 'april', 'may', 'june', 'july',
                                       'august', 'september', 'october', 'november', 'december'], 1)}
YEAR = 2026
DAYS = ['Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday']
# A day the class is not in its room.
AWAY = re.compile(r'holiday|vacation|chaturthi|chaturdashi|jayanti|dasara|dussehra|diwali|christmas|independence'
                  r'|exam|prep leave|orientation|induction|sod event|study tour|overseas exchange|industry\s*exposure', re.I)
# The tab's programme, from its name.
PROGRAMME = [(r'_TD|\bTD_', 'Transportation Design'), (r'UIUX', 'UI/UX Design'), (r'Fashion', 'Fashion Design'),
             (r'Interior', 'Interior Design'), (r'Product', 'Product Design'), (r'VCD', 'Visual Communication Design'),
             (r'Foundation', 'Foundation')]
STOP = {'of', 'and', 'the', 'for', 'in', 'to', 'a', 'an', 'with'}

tidy = lambda s: re.sub(r'\s+', ' ', str(s or '')).strip()
norm = lambda s: re.sub(r'[^a-z0-9]+', ' ', s.lower()).strip()


def initials(name, stop=True):
    return ''.join(w[0] for w in re.findall(r'[A-Za-z0-9]+', name) if not stop or w.lower() not in STOP).upper()


ROMAN = {1: 'I', 2: 'II', 3: 'III', 4: 'IV', 5: 'V', 6: 'VI'}
# Codes no rule finds: "TW" is Two Wheeler, and AUX takes its X from "UX".
ALIAS = {'TWSR': '2 Wheeler Sketching', 'AUX': 'Automotive UX'}
# A programme whose sheet names no classroom: the room its classes are put in.
ROOMLESS = {'Interior Design': 'Interior Design Studio'}
# Courses the team teaches together, not the lead alone: the dean is on
# Studio-II's team, and only the team says so.
TEAM_TAUGHT = {'Interior Design Studio-II'}
# Cell suffixes that are still the course, in its room.
SUFFIX = re.compile(r'\s*(internal\s+presentati\w*|\.?\s*practice sessions)\s*$', re.I)


def teacher(lead, team, whole_team=False):
    """The course lead, else its team: 'Name (Game Design)' -> 'Name', 'A and B' -> 'A / B'.
    whole_team: the lead and the team, each person once."""
    who = re.sub(r'\(.*?\)', '', f'{lead} / {team}' if whole_team else lead or team)
    who = re.sub(r'\s*(/|,|&|\band\b)\s*', ' / ', who)
    parts, seen = [], set()
    for w in (tidy(w) for w in who.split(' / ')):
        key = re.sub(r'^(prof|dr|ar)\b\.?\s*', '', w.lower())
        if w and w.lower() not in ('new faculty', 'vf') and key not in seen:
            seen.add(key)
            parts.append(w)
    return ' / '.join(parts)


def course_name(s):
    s = re.sub(r'^[A-Z]{1,2}\d{3,4}[A-Z]?\s*-\s*', '', s)     # "D0310B- Design Studio 1", "TD409D- ..."
    s = re.sub(r'\(\s*common for all\s*\)', '', s, flags=re.I)
    return tidy(s.split('\n')[0]).rstrip(' -:')


def room_list(header):
    m = re.search(r'Class Room(?: No\.?)?\s*:?\s*(.*)', header, re.I)
    text = tidy(m.group(1)) if m else ''
    lab = re.match(r'(\d{3})\s*-\s*(.+?)_', text)                # "107-Cintiq Lab_First Floor"
    if lab: return [f'{lab.group(1)} ({lab.group(2)})']
    return re.findall(r'\b\d{3}\b', re.sub(r'\(.*?\)', '', text))


wb = openpyxl.load_workbook(sys.argv[1])
out = {'source': sys.argv[1].replace('\\', '/').split('/')[-1], 'building': 'School of Design',
       'note': "Converted by convert-design.py from each tab's semester calendar: one subject per session, 9.30 to 12.30 and 1.00 to 4.00, in the room the header names, kept as dates. Holidays, exams, events and days away are left out.",
       'classes': [], 'calendars': []}
seen = {}
for ws in wb.worksheets:
    tab = ws.title.strip()
    merged = {}
    for m in ws.merged_cells.ranges:
        merged[(m.min_row, m.min_col)] = (m.max_row, m.max_col)
    cell = lambda r, c: tidy(ws.cell(r, c).value)
    header = ' '.join(cell(r, c) for r in range(1, 6) for c in range(1, ws.max_column + 1))
    sem = re.search(r'Semester\s*:\s*([IVX]+)\b', header)
    prog = next((p for pat, p in PROGRAMME if re.search(pat, tab, re.I)), None)
    level = 'M.Des' if tab.upper().startswith('M.DES') else 'B.Des'
    if not sem or not prog:
        print(f'{tab}: no semester or programme, skipped'); continue
    label = f'{level} {prog} Sem {sem.group(1)}'
    rooms = room_list(header) or ([ROOMLESS[prog]] if prog in ROOMLESS else [])
    if not rooms:
        print(f'{label}: the sheet names no classroom, left out'); continue

    head = next(r for r in range(1, 12) if any(cell(r, c).upper() == 'MONDAY' for c in range(1, ws.max_column + 1)))
    cols = {}   # day => date column; the subject is the column after it
    for c in range(1, ws.max_column + 1):
        if cell(head, c).capitalize() in DAYS: cols[cell(head, c).capitalize()] = c
    course_row = next(r for r in range(head, ws.max_row + 1) if cell(r, 1) == 'Course')

    # The course list's own columns: they move between workbooks.
    hdr = {cell(course_row, c): c for c in range(1, ws.max_column + 1)}
    courses = []
    for r in range(course_row + 1, ws.max_row + 1):
        name = course_name(str(ws.cell(r, 1).value or ''))
        if not name or name.startswith('Prepared'): break
        courses.append({'name': name, 'who': teacher(cell(r, hdr['Course Lead']), cell(r, hdr['Course Team']), name in TEAM_TAUGHT),
                        'ini': (initials(name), initials(name, False))})

    def find(text):
        """The course a cell names, or None."""
        text = SUFFIX.sub('', text)
        if text in ALIAS:
            return next((c for c in courses if c['name'].startswith(ALIAS[text])), None)
        n = norm(text)
        for c in sorted(courses, key=lambda c: -len(c['name'])):
            cn = norm(c['name'])
            if cn and (n == cn or n.startswith(cn) or cn.startswith(n) and len(n) > 6): return c
        # "DAF", "DS4- II", "ESI: Module 01", "ESI Module 05": a course by its initials
        ab = re.match(r'^([A-Z]{2,5})\s*-?\s*(\d?)\s*(?:-\s*(I{1,3}))?(?=$|:|\s+Module)', text)
        if ab:
            letters, digit, part = ab.groups()
            hits = [c for c in courses if any(i.startswith(letters) for i in c['ini'])
                    and (not digit or re.search(rf'\b({digit}|{ROMAN[int(digit)]})\b', c['name']))
                    and (not part or re.search(rf'-\s*{part}$', c['name']))]
            if not part and len(hits) > 1: hits = [c for c in hits if not re.search(r'-\s*I+$', c['name'])]
            if len(hits) == 1: return hits[0]
        # "Craft Exp. & Creative Artwork: Module 01": the words a course starts with
        words = [w for w in norm(text.split(':')[0]).split() if w not in STOP]
        for c in courses:
            cw = [w for w in norm(c['name']).split() if w not in STOP]
            if len(words) >= 2 and all(any(x.startswith(w[:4]) for x in cw) for w in words[:3]): return c
        return None

    # The Monday each week row stands for, checked against the dates the sheet prints.
    subjects = []   # [name, teacher], indexed by the dates
    dates = {}      # date => [morning, afternoon], a subject's index or -1
    cells = {}      # date => [[first row, first col, last row, last col] or None, ...], 0-based
    month = None
    monday = None
    for r in range(head + 1, course_row):
        if cell(r, 3) != 'Mor.' or not cell(r, 1).isdigit(): continue
        if cell(r, 2).lower() in MONTHS: month = MONTHS[cell(r, 2).lower()]
        num = re.match(r'[0-9]+', cell(r, cols['Monday']))
        if monday is None:
            monday = datetime.date(YEAR, month, int(num.group()))
        else:
            monday += datetime.timedelta(days=7)
            if num and int(num.group()) != monday.day: print(f'{label} row {r}: Monday says {num.group()}, expected {monday}')
        for i, day in enumerate(DAYS):
            when = (monday + datetime.timedelta(days=i)).isoformat()
            c = cols[day] + 1
            for k, row in enumerate((r, r + 1)):
                text = cell(row, c)
                box = (row, c) + merged.get((row, c), (row, c))
                if k == 1 and not text and merged.get((r, c), (r, c))[0] > r:
                    # merged over both rows: the morning's subject all day
                    text, box = cell(r, c), (r, c) + merged[(r, c)]
                if not text or text.isdigit(): continue
                if AWAY.search(text):
                    seen.setdefault(('away', text), label); continue
                course = find(text)
                subject = [course['name'] if course else text, course['who'] if course else '']
                seen.setdefault(('course' if course else 'plain', f'{text}  ->  {subject[0]}' + (f'  [{subject[1]}]' if subject[1] else '')), label)
                if subject not in subjects: subjects.append(subject)
                dates.setdefault(when, [-1, -1])[k] = subjects.index(subject)
                cells.setdefault(when, [None, None])[k] = [box[0] - 1, box[1] - 1, box[2] - 1, box[3] - 1]

    out['calendars'].append({'class': label, 'page': tab, 'rooms': rooms, 'sessions': [['9.30', 3], ['1.00', 3]],
                             'subjects': subjects, 'dates': dates, 'cells': cells})
    print(f'{label}: {len(dates)} days in {", ".join(rooms)}, {len(subjects)} subjects ({len(courses)} courses listed)')

for (kind, text), label in sorted(seen.items()):
    print(f'  {kind:6} {text}   ({label})')
json.dump(out, open(sys.argv[2], 'w', encoding='utf-8'), indent=1, ensure_ascii=False)
