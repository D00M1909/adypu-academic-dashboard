# Dumps the class tabs of a "Time / Day" timetable workbook (one tab per
# division, a time column, Monday to Friday across, merged cells for long
# periods) as JSON for convert-cse.php: each period's rows, each cell's text
# and the rows its merge covers, and the course list under the grid.
#
#   python dump-grid.py "SY CSE-TimeTable.xlsx" sy-cse.grid.json
#
# Needs openpyxl, which the dashboard's own xlsx reader does not replace: it
# does not see merges.
import json, re, sys
import openpyxl

wb = openpyxl.load_workbook(sys.argv[1])
out = []
for ws in wb.worksheets:
    merges = {(m.min_row, m.min_col): (m.max_row, m.max_col) for m in ws.merged_cells.ranges}
    covered = {}
    for m in ws.merged_cells.ranges:
        for r in range(m.min_row, m.max_row + 1):
            for c in range(m.min_col, m.max_col + 1):
                covered[(r, c)] = (m.min_row, m.min_col)
    head = None
    for r in range(1, 15):
        for c in range(1, 4):
            if re.match(r'Time\s*/\s*Day', str(ws.cell(r, c).value or '')):
                head, tcol = r, c
    if head is None:
        continue
    days = []
    for c in range(tcol + 1, ws.max_column + 1):
        v = str(ws.cell(head, c).value or '').strip()
        if v in ('Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday'):
            last = merges.get((head, c), (head, c))[1]
            days.append({'day': v, 'cols': [c, last]})
    periods, cells, courses = [], [], []
    stop = None
    for r in range(head + 1, ws.max_row + 1):
        label = str(ws.cell(r, tcol).value or '').strip()
        if 'Course Name' in ' '.join(str(ws.cell(r, c).value or '') for c in range(1, ws.max_column + 1)):
            stop = r
            break
        m = re.match(r'(\d{1,2})\.(\d{2})\s*to\s*(\d{1,2})\.(\d{2})', label)
        if m:
            periods.append({'label': f'{int(m.group(1))}.{m.group(2)} to {int(m.group(3))}.{m.group(4)}', 'rows': [r]})
        elif periods:
            periods[-1]['rows'].append(r)
    for p in periods:
        for r in p['rows']:
            for d in days:
                for c in range(d['cols'][0], d['cols'][1] + 1):
                    if covered.get((r, c), (r, c)) != (r, c): continue
                    v = str(ws.cell(r, c).value or '').strip()
                    if not v: continue
                    lr, lc = merges.get((r, c), (r, c))
                    cells.append({'day': d['day'], 'row': r, 'last_row': lr, 'col': c, 'text': re.sub(r'\s*\n\s*', ' | ', v)})
    if stop:
        hdr = [str(ws.cell(stop, c).value or '').strip() for c in range(1, ws.max_column + 1)]
        ncol = hdr.index('Course Name') + 1
        fcol = hdr.index('Faculty Name') + 1
        for r in range(stop + 1, ws.max_row + 1):
            n, f = str(ws.cell(r, ncol).value or '').strip(), str(ws.cell(r, fcol).value or '').strip()
            if n.startswith('Prepared') or str(ws.cell(r, 1).value or '').startswith('Prepared'): break
            if n: courses.append({'course': n, 'faculty': re.sub(r'\s*\n\s*', ' ', f)})
    out.append({'tab': ws.title, 'periods': periods, 'cells': cells, 'courses': courses})
json.dump(out, open(sys.argv[2], 'w', encoding='utf-8'), indent=1, ensure_ascii=False)
print(sys.argv[2], [(t['tab'], len(t['periods']), len(t['cells']), len(t['courses'])) for t in out])
