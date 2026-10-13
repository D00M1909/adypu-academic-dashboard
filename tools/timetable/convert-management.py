# Converts the School of Management timetable workbook into the format
# build.php reads.
#
#   python convert-management.py "Timetable School of Mgt.xlsx" management-2026-10-12.json
#
# Needs openpyxl. Each tab is one programme: Monday to Friday across, one row
# per hour, and a course list under it (code, name, faculty, specialisation).
# A cell names one course, or several run in parallel (BBA Sem V and MBA Sem
# III specialisations), separated by "---", "___" or only a blank line, so
# courses are found by matching the tab's own course names, longest first,
# rather than by splitting. Rooms come from the header line: one room for the
# tab, or one per specialisation. A course whose specialisation has no room
# (a common course, Healthcare) is printed and left out.
import json, re, sys
import openpyxl

DAYS = {'B': 'Monday', 'C': 'Tuesday', 'D': 'Wednesday', 'E': 'Thursday', 'F': 'Friday'}
TABS = {
    'BBA-I':   ('BBA Sem I',   {'*': '506A'}),
    'BBA-III': ('BBA Sem III', {'*': '603'}),
    'BBA-V':   ('BBA Sem V',   {'Finance': '601', 'Entrepreneurship': '602', 'Marketing': '604', 'FinTech': '605'}),
    # 701, except the 12.30 period, which the header puts in 606
    'MBA-I':   ('MBA Sem I',   {'*': '701', '12:30': '606'}),
    'MBA-III': ('MBA Sem III', {'Marketing': '609', 'Finance': '606', 'Human Resource': '712', 'Fintech': 'IT Lab (S2)',
                                'Entrepreneurship': '610', 'Operations': '805'}),
}
norm = lambda s: re.sub(r'[^a-z0-9&]+', ' ', s.lower()).strip()

wb = openpyxl.load_workbook(sys.argv[1])
out = {'source': 'Timetable School of Mgt (1) (2).xlsx', 'building': 'ULC 1',
       'note': 'Converted by convert-management.py. Rooms from each tab\'s header line (one room, or one per specialisation). Common courses and Healthcare name no room and are left out. Building not stated in the sheet: taken as ULC 1 (603 and 805 are known ULC 1 rooms).',
       'faculty': {'': ''}, 'subjects': {}, 'classes': []}
for tab, (label, rooms) in TABS.items():
    ws = wb[tab]
    head = next(r for r in range(1, 12) if str(ws.cell(r, 1).value or '').strip() == 'Time')
    listed = next(r for r in range(head, ws.max_row + 1) if str(ws.cell(r, 1).value or '').strip() == 'Course Code')
    courses = []
    for r in range(listed + 1, ws.max_row + 1):
        code, name, fac = (str(ws.cell(r, c).value or '').strip() for c in (1, 2, 4))
        spec = str(ws.cell(r, 6).value or '').strip()
        if not code or not name: continue
        short = re.sub(r'-\s*(Minor|Common Course|Vocational Course)$', '', name).strip()
        courses.append({'code': code, 'name': short, 'fac': fac, 'spec': spec})
    by_len = sorted(courses, key=lambda c: -len(c['name']))
    per_spec = {}
    for r in range(head + 1, listed):
        time = str(ws.cell(r, 1).value or '').strip()
        m = re.match(r'(\d{2}):(\d{2})-(\d{2}):(\d{2})', time)
        if not m: continue
        start = f"{int(m.group(1))}.{m.group(2)}"
        for col, day in DAYS.items():
            text = ' ' + norm(str(ws[f'{col}{r}'].value or '')) + ' '
            if not text.strip(): continue
            found = []
            for c in by_len:
                n = ' ' + norm(c['name']) + ' '
                while n in text:
                    found.append(c)
                    text = text.replace(n, ' # ', 1)
            left = re.sub(r'\b(b|x)\d{4,5}[ab]\b|#|dr|mr|ms|prof|library period', ' ', text)
            for c in found:
                if '*' in rooms:
                    room = rooms.get(m.group(1) + ':' + m.group(2), rooms['*'])
                    group = label
                else:
                    room = rooms.get(c['spec'])
                    group = f"{label} {c['spec']}"
                if room is None:
                    print(f'  {tab} {day} {start}: {c["name"]} ({c["spec"] or "no specialisation"}) has no room, left out')
                    continue
                key = f'{tab}|{c["code"]}'
                out['subjects'][key] = c['name']
                out['faculty'][c['fac']] = c['fac']
                per_spec.setdefault(group, []).append([day, start, 1, key, room, c['fac'], False, ''])
            if not found and 'library' not in text:
                print(f'  {tab} {day} {start}: nothing recognised in "{ws[f"{col}{r}"].value}"')
    for group, entries in per_spec.items():
        out['classes'].append({'class': group, 'page': tab, 'entries': entries})
        print(f'{group}: {len(entries)} bookings')
json.dump(out, open(sys.argv[2], 'w', encoding='utf-8'), indent=1, ensure_ascii=False)
